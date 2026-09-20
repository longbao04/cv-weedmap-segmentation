#!/usr/bin/env python3
"""Regenerate the existing supervisor-facing WeedMap progress report."""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "reports" / "weedmap_progress_report.docx"
INK, NAVY, PALE, GRID = "17212B", "213B58", "F2F6FA", "D9D9D9"


def set_font(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(INK)
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for name in ("ascii", "hAnsi"):
        fonts.set(qn(f"w:{name}"), "Arial")
    fonts.set(qn("w:eastAsia"), "PingFang SC")


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(1.9)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    styles = doc.styles
    set_font(styles["Normal"], 10)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    styles["Normal"].paragraph_format.space_after = Pt(5)
    set_font(styles["Title"], 20, True)
    styles["Title"].paragraph_format.space_after = Pt(12)
    ppr = styles["Title"]._element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    for name, size, before, after in (("Heading 1", 13, 15, 6), ("Heading 2", 10.5, 9, 4)):
        set_font(styles[name], size, True)
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("WeedMap 实验进度报告")


def paragraph(doc, text):
    return doc.add_paragraph(text)


def section(doc, number, title):
    doc.add_heading(f"{number} {title}", level=1)


def shade(cell, color):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)


def table(doc, headers, rows, widths, size=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, width in enumerate(widths):
        t.columns[i].width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "5")
        edge.set(qn("w:color"), GRID)
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate([headers, *rows]):
        row = t.rows[0] if row_index == 0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.width = Cm(widths[i])
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index == 0:
                shade(cell, NAVY)
            elif row_index % 2 == 0:
                shade(cell, PALE)
            mar = OxmlElement("w:tcMar")
            for side, amount in (("top", 95), ("bottom", 95), ("start", 90), ("end", 90)):
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:w"), str(amount))
                item.set(qn("w:type"), "dxa")
                mar.append(item)
            cell._tc.get_or_add_tcPr().append(mar)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else INK)
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def metrics(doc, rows, size=7.5):
    table(doc, ["模型 / 损失", "Pixel Acc", "Mean IoU", "Background IoU", "Crop IoU", "Weed IoU"],
          rows, [4.2, 2.35, 2.35, 2.8, 2.5, 2.8], size)


def build():
    doc = Document()
    setup(doc)
    doc.core_properties.title = "WeedMap 多光谱语义分割实验进度报告"
    doc.add_paragraph("WeedMap 多光谱语义分割实验进度报告", style="Title")
    paragraph(doc, "面向导师的阶段性汇报。报告总结真实 WeedMap 数据处理、统一样本对照、三 seed 语义分割实验、VCR regression attention router，以及 YOLO 检测数据构建、后处理阈值和 MobileNetV3-Small-YOLOv8 轻量化实验。当前语义分割主线最佳设置为 MobileNetV2ShallowUNet 方案 A + boundary weighted CE r5_w4：验证集 Pixel accuracy 96.98% ± 0.24%，Mean IoU 76.71% ± 0.55%，Weed IoU 59.82% ± 0.64%。YOLOv8n crop+weed r010 仍是检测基线；MobileNetV3 变体减小了参数量和 GFLOPs，但精度及当前 CPU 延迟尚未超过该基线。")

    section(doc, 1, "项目目标与数据集")
    paragraph(doc, "项目目标是利用无人机 RGB 与多光谱影像，将每个有效像素分为 background、crop、weed，为田间杂草识别和精准除草提供区域信息。前期通过 NDVI、NDRE 模拟光谱差异，并用 synthetic 数据跑通分割流程；报告中的主要结论依据真实 WeedMap Tiles 实验。")
    paragraph(doc, "真实 tile 尺寸为 360 × 480。原始数据检查得到 1670 个候选样本；后续严格对照使用 RGB 与多光谱均可用的 WeedMap common split，共 454 张，固定划分为训练 363 张、验证 91 张。多光谱输入为 G、R、RedEdge、NIR、NDVI 五通道。")

    section(doc, 2, "真实 WeedMap 数据处理与标签映射")
    paragraph(doc, "已检查影像、GroundTruth 与有效区 mask 的空间对应关系。本地数据中 GroundTruth_color 的类别语义最明确，Dataset 从 color 标签生成 0/1/2 类别；mask=255 标记无效区域，在训练和评价中作为 ignore_index。")
    table(doc, ["数据", "原值", "训练含义"], [
        ("GroundTruth_color", "black / green / red", "background=0 / crop=1 / weed=2"),
        ("GroundTruth_iMap", "0 / 10000 / 2", "对应 background / crop / weed；10000 不是 ignore"),
        ("mask", "0 / 255", "有效区域 / 无效区域"),
    ], [4, 4.7, 8.3], 8.8)

    section(doc, 3, "RGB 与 multispectral 公平对比")
    paragraph(doc, "两组使用同一 454 张样本、相同 363/91 划分、SmallUNet、Weighted CE、类别权重 1/4/8、10 epochs 和 batch size 2。下表是验证集最后一轮结果；这是输入通道对照，不能与后文 20 epochs 的 best checkpoint 直接混为同一设置。")
    metrics(doc, [
        ("RGB", "94.28%", "65.27%", "95.49%", "65.75%", "34.56%"),
        ("Multispectral", "94.38%", "67.76%", "94.94%", "61.62%", "46.73%"),
    ], 8)
    paragraph(doc, "多光谱 Mean IoU 高 2.49 个百分点，Weed IoU 高 12.17 个百分点；RGB Crop IoU 高 4.13 个百分点。该结果支持后续以多光谱作为主输入，并保留对 crop 类差异的关注。")

    section(doc, 4, "SmallUNet baseline")
    paragraph(doc, "真实数据主 baseline：WeedMap common split，多光谱输入，SmallUNet，Weighted CE，类别权重 background/crop/weed=1/4/8，20 epochs，batch size 2，seed 0/1/2。每个 seed 按验证集 Mean IoU 选择 best checkpoint；下表的 ± 为三 seed 样本标准差。")
    metrics(doc, [("SmallUNet + Weighted CE", "96.38% ± 0.36%", "74.78% ± 0.82%", "96.70% ± 0.43%", "70.76% ± 1.11%", "56.89% ± 2.06%")])
    paragraph(doc, "三个 seed 的 best epoch 均为 15。早期 synthetic 普通 CE 的 Weed IoU 仅 0.26%，加入 class weights 后 5 epochs 达到 75.54%；这一模拟实验说明少数类权重的重要性，但数值不代表真实田间性能。")

    section(doc, 5, "Loss 对比：Weighted CE / Focal Loss / Dice + CE")
    paragraph(doc, "固定 common split、多光谱 SmallUNet、20 epochs、batch size 2、seed 0/1/2 和 best checkpoint 选择规则，仅比较损失函数。")
    metrics(doc, [
        ("Weighted CE", "96.38% ± 0.36%", "74.78% ± 0.82%", "96.70% ± 0.43%", "70.76% ± 1.11%", "56.89% ± 2.06%"),
        ("Focal Loss", "96.69% ± 0.26%", "74.00% ± 1.25%", "97.19% ± 0.11%", "69.27% ± 3.00%", "55.55% ± 1.72%"),
        ("Dice + CE", "96.73% ± 0.24%", "74.45% ± 3.37%", "97.21% ± 0.05%", "69.54% ± 3.82%", "56.61% ± 6.31%"),
    ])
    paragraph(doc, "Weighted CE 的平均 Mean IoU 和 Weed IoU 均最高，作为稳定 baseline。Dice + CE 的 seed 0/1 表现较强，但 seed 2 明显下降，跨 seed 波动较大；Focal Loss 的平均 Weed IoU 略低。")

    section(doc, 6, "Boundary error analysis")
    paragraph(doc, "对 SmallUNet + Weighted CE 的 20 epochs best checkpoint 分析 sample index=0，排除 label=255 后共 112158 个有效像素、5812 个错误像素。5px 边界区域含 38199 个有效像素，占有效像素 34.06%，却包含 5532 个错误，占全部错误 95.18%。")
    table(doc, ["区域", "错误像素", "区域错误率"], [
        ("5px 边界内", "5532", "14.48%"),
        ("5px 边界外", "280", "0.38%"),
    ], [7, 4.5, 5.5], 8.8)
    paragraph(doc, "该单张样本表明错误集中在边界，因而尝试对边界像素增加损失权重；它不能单独代表整体验证集。")

    section(doc, 7, "SmallUNet + boundary weighted CE r5_w4")
    paragraph(doc, "在 SmallUNet baseline 上使用 boundary_weighted_ce，边界半径 5px、边界权重 4.0，类别权重仍为 1/4/8。common split、多光谱、20 epochs、batch size 2、三 seed 与 best checkpoint 规则保持一致。")
    metrics(doc, [("SmallUNet + boundary CE r5_w4", "96.79% ± 0.16%", "75.17% ± 1.02%", "97.15% ± 0.11%", "70.07% ± 2.41%", "58.29% ± 0.92%")])
    paragraph(doc, "相对 SmallUNet + Weighted CE，Mean IoU 提高 0.39 个百分点，Weed IoU 提高 1.40 个百分点；Crop IoU 均值下降 0.69 个百分点。sample index=0 的 5px 边界错误率由 14.48% 降至 9.50%，总错误像素由 5812 降至 3746。")

    section(doc, 8, "MobileNetV2ShallowUNet 方案 A")
    paragraph(doc, "方案 A 是 MobileNetV2-style shallow encoder 改造实验，不是论文完整 MobileNetV2-U-Net 复现。它保留当前 SmallUNet 的 C1/C2、两级 Decoder、skip connection 和 segmentation head，以浅层 MobileNetV2-style inverted residual blocks B1-B6 替换原 enc2 与 bottleneck。完整 B1-B17 多尺度结构留作方案 B。")
    paragraph(doc, "参数量由 SmallUNet 的 117363 降至 107155，减少约 8.7%。在相同 common split、多光谱、Weighted CE、20 epochs、batch size 2、三 seed 条件下：")
    metrics(doc, [("方案 A + Weighted CE", "96.53% ± 0.47%", "75.86% ± 1.22%", "96.80% ± 0.44%", "71.74% ± 2.76%", "59.03% ± 1.74%")])
    paragraph(doc, "与原 SmallUNet + Weighted CE 比较，Mean IoU 高 1.08 个百分点，Weed IoU 高 2.14 个百分点。")

    section(doc, 9, "MobileNetV2ShallowUNet + boundary weighted CE r5_w4")
    paragraph(doc, "当前最强语义分割设置：WeedMap common split；multispectral；MobileNetV2ShallowUNet 方案 A；boundary_weighted_ce；boundary radius=5、boundary weight=4.0；类别权重 background=1.0、crop=4.0、weed=8.0；20 epochs、batch size=2、seed=0/1/2；每个 seed 按验证集 Mean IoU 选择 best checkpoint。")
    table(doc, ["Seed", "Best epoch", "Pixel Acc", "Mean IoU", "Background IoU", "Crop IoU", "Weed IoU"], [
        ("0", "20", "96.91%", "76.97%", "97.20%", "73.41%", "60.30%"),
        ("1", "12", "96.78%", "76.08%", "97.15%", "71.02%", "60.07%"),
        ("2", "18", "97.25%", "77.07%", "97.44%", "74.69%", "59.10%"),
        ("均值 ± 标准差", "—", "96.98% ± 0.24%", "76.71% ± 0.55%", "97.26% ± 0.16%", "73.04% ± 1.86%", "59.82% ± 0.64%"),
    ], [2.5, 2, 2.35, 2.35, 2.8, 2.5, 2.5], 7.4)
    doc.add_heading("与前序主线对比", level=2)
    metrics(doc, [
        ("SmallUNet + Weighted CE", "96.38% ± 0.36%", "74.78% ± 0.82%", "96.70% ± 0.43%", "70.76% ± 1.11%", "56.89% ± 2.06%"),
        ("SmallUNet + boundary CE r5_w4", "96.79% ± 0.16%", "75.17% ± 1.02%", "97.15% ± 0.11%", "70.07% ± 2.41%", "58.29% ± 0.92%"),
        ("方案 A + Weighted CE", "96.53% ± 0.47%", "75.86% ± 1.22%", "96.80% ± 0.44%", "71.74% ± 2.76%", "59.03% ± 1.74%"),
        ("方案 A + boundary CE r5_w4", "96.98% ± 0.24%", "76.71% ± 0.55%", "97.26% ± 0.16%", "73.04% ± 1.86%", "59.82% ± 0.64%"),
    ], 7.4)
    paragraph(doc, "相对 SmallUNet + Weighted CE baseline，Mean IoU 从 74.78% ± 0.82% 提升至 76.71% ± 0.55%（+1.93 个百分点）；Weed IoU 从 56.89% ± 2.06% 提升至 59.82% ± 0.64%（+2.93 个百分点）。结果表明方案 A 的浅层 encoder 改造与 boundary weighted CE 在当前实验中可以叠加提升，尤其对 weed 类识别更有帮助。")

    section(doc, 10, "VCR 植被覆盖率评判指标")
    paragraph(doc, "VCR = vegetation pixels / valid pixels。其中 vegetation pixels 指 NDVI > 0.2 且 label != 255 的像素；valid pixels 指 label != 255 的像素。VCR 用于场景划分，不参与分割模型训练或 IoU 计算。")
    table(doc, ["场景", "初始 VCR 阈值", "样本数", "比例"], [
        ("sparse", "VCR < 0.20", "13 / 454", "≈ 2.86%"),
        ("transition", "0.20 ≤ VCR ≤ 0.30", "9 / 454", "≈ 1.98%"),
        ("dense", "VCR > 0.30", "432 / 454", "≈ 95.15%"),
    ], [3.5, 6, 3.5, 4], 8.8)
    paragraph(doc, "WeedMap common split 共 454 张样本：sparse 13 张、transition 9 张、dense 432 张；VCR 均值 0.7928、中位数 0.9116、最小值 0.0000、最大值 0.9999。dense 占 432 / 454 ≈ 95.15%，sparse 占 13 / 454 ≈ 2.86%，transition 占 9 / 454 ≈ 1.98%。当前绝大多数样本属于高植被覆盖场景，因此继续优化 U-Net / MobileNetV2ShallowUNet 语义分割模型是合理的；YOLO 更适合作为低覆盖稀疏场景下的候选路线或后续扩展。0.20 和 0.30 是初始经验阈值。")

    section(doc, 11, "VCR Regression Attention Router")
    paragraph(doc, "目的：在 weed recognition pipeline 前端，根据场景植被覆盖程度选择后续分支。模型不直接做三分类，而是先从原始 G、R、RE、NIR 四个波段回归连续 VCR，再按 predicted VCR < 0.20、0.20 至 0.30、> 0.30 分别路由到 sparse YOLO、transition 双模型或人工确认、dense semantic segmentation。不输入 NDVI，因为 VCR 标签由 NDVI threshold 生成，直接输入 NDVI 可能只是复现标签规则。")
    paragraph(doc, "数据准备使用现有 VCR summary，并生成 outputs/vcr_regression_samples.csv；454 条样本全部校验通过。由于 sparse、transition、dense 仅有 13、9、432 张，直接三分类容易全部预测 dense 并获得虚高 accuracy，因此采用连续 VCR regression。")
    table(doc, ["模型", "正式设置", "关键结果"], [
        ("Tiny CNN", "30 epochs × 3 seeds", "MAE 0.1014 ± 0.0101"),
        ("SE-Tiny CNN", "30 epochs × 3 seeds", "MAE 0.1055 ± 0.0044"),
        ("CBAM-Tiny CNN", "30 epochs × 3 seeds", "MAE 0.0923 ± 0.0133；sparse recall 0.1111 ± 0.1925"),
    ], [4.0, 5.0, 8.0], 8.2)
    paragraph(doc, "正式实验实际使用 Huber loss（delta=0.1），各 run 按 validation MAE 选择 best checkpoint。CBAM-Tiny CNN 的平均 MAE 最低，说明 attention 可能有助于学习植被覆盖程度；但 non-dense 样本太少，validation sparse 仅 3 张，每错一张都会造成很大波动，且特定 subset 与相邻帧可能带来空间数据泄漏风险。")
    paragraph(doc, "结论：VCR regression attention router 的探索性实验有效完成，但 sparse recall 仍然不足，当前不能替代 NDVI-VCR 零训练规则，也不能作为部署模块。项目主线仍是 dense vegetation 场景下的 MobileNetV2ShallowUNet + boundary weighted CE r5_w4 语义分割；YOLO 是 sparse 场景候选路线，VCR router 的长期作用是连接两条分支。后续需增加 sparse / transition 样本，并采用 repeated stratified cross-validation 和 subset-level holdout。")

    section(doc, 12, "YOLO detection pipeline 的两个阶段")
    doc.add_heading("A Detection Dataset 构建阶段", level=2)
    paragraph(doc, "WeedMap semantic mask → target mask → connected components → component bbox → bbox statistics → visualization / statistical analysis → dataset bbox filtering rules → YOLO labels → Detection Dataset。")
    paragraph(doc, "WeedMap 原始标签没有 instance identity。connected component 只能近似生成 bbox：稀疏场景中的独立小 component 可能接近单株或单簇；密集粘连区域的大 component 可能包含多株植株或片状杂草，不能直接解释为单株。这一阶段的 filtering 是训练标签构建阶段的 bbox 清洗，应结合 width、height、area、aspect ratio 和 bbox/image area ratio 分布及框叠加可视化确定规则，避免误删真实小 weed。")
    doc.add_heading("B YOLO 训练与推理阶段", level=2)
    paragraph(doc, "Detection Dataset → YOLOv8n smoke test / training → network candidate predictions → confidence filtering → NMS → final bounding boxes。")
    paragraph(doc, "YOLO 网络输出候选框、类别和置信度；confidence filtering 去掉低置信度预测框，NMS 去掉高度重叠的重复预测框。两者属于 inference post-processing，不属于 backbone，也不同于 A 阶段用于生成训练标签的 dataset bbox filtering。当前 YOLOv8n 验证的是 mask → connected components → bbox → detection dataset → training/inference 流程，不是 MobileNetV3-YOLOv3 复现。")

    section(doc, 13, "YOLO bbox 分布与标签局限")
    paragraph(doc, "当前 crop+weed YOLO baseline 检测数据集的 bbox 统计如下。YOLO labels 由 WeedMap semantic mask 的 target mask 和 connected components 自动生成，不是人工标注的 instance bbox。")
    table(doc, ["统计项", "结果"], [
        ("Total boxes", "19,443"),
        ("Crop boxes", "12,311"),
        ("Weed boxes", "7,132"),
        ("Mean boxes per image", "42.83"),
        ("Min / median / max boxes per image", "0 / 44 / 120"),
    ], [9, 8], 8.8)
    paragraph(doc, "统计由 analyze_yolo_bbox_statistics.py 生成，用于评估 min-area、max-area-ratio、min-width 和 min-height 等规则。连通域框只是实例目标的近似标签；bbox 分布统计与可视化用于判断数据集阶段的过滤规则，不会自动修改标签，也不能替代人工 instance annotations。")

    section(doc, 14, "YOLOv8n bbox filtering 对比")
    paragraph(doc, "r020、r015、r010 均使用 YOLOv8n、5 epochs、imgsz=480、batch=4、device=mps、seed=0，以及相同的 363/91 train/validation images。仅调整数据集构建阶段的 max-box-area-ratio；min-area=80、min-box-width=6、min-box-height=6 保持不变。")
    table(doc, ["设置", "类别", "P", "R", "mAP50", "mAP50-95"], [
        ("r020", "all", "0.498", "0.594", "0.526", "0.291"),
        ("r020", "crop", "0.525", "0.650", "0.595", "0.370"),
        ("r020", "weed", "0.470", "0.538", "0.457", "0.212"),
        ("r015", "all", "0.508", "0.587", "0.530", "0.297"),
        ("r015", "crop", "0.522", "0.672", "0.616", "0.387"),
        ("r015", "weed", "0.492", "0.505", "0.445", "0.206"),
        ("r010", "all", "0.505", "0.591", "0.531", "0.296"),
        ("r010", "crop", "0.527", "0.652", "0.601", "0.376"),
        ("r010", "weed", "0.482", "0.530", "0.460", "0.215"),
    ], [2.4, 2.4, 2.5, 2.5, 3.1, 4.0], 7.5)
    paragraph(doc, "相对 r020，r010 的 weed mAP50 和 mAP50-95 均提高 0.003，weed recall 下降 0.008。r015 的 all mAP50 与 r010 几乎持平、mAP50-95 略高，但 weed recall 和 weed mAP 均低于 r010。综合整体与 weed 指标，r010 保留为当前最佳过滤设置；三组差异很小，不能声称过滤阈值带来显著提升。")

    section(doc, 15, "YOLO weed-only 与 crop+weed 对照")
    paragraph(doc, "在 r010 数据集过滤设置下比较两类检测与只检测 weed。下表均为验证集 weed 类指标。")
    table(doc, ["设置", "Weed P", "Weed R", "Weed mAP50", "Weed mAP50-95"], [
        ("crop+weed r010", "0.482", "0.530", "0.460", "0.215"),
        ("weed-only r010", "0.431", "0.488", "0.421", "0.182"),
    ], [4.6, 2.5, 2.5, 3.6, 3.8], 8.2)
    paragraph(doc, "当前数据构建方式、YOLOv8n、5 epochs、seed=0 条件下，weed-only 的四项 weed 指标均未超过 crop+weed。保留 crop 类可能有助于模型学习作物行结构以及 weed/crop 区分，但这不构成对 weed-only 路线的最终否定；crop+weed r010 作为主要检测设置，weed-only r010 作为对照记录。")

    section(doc, 16, "YOLO 预测可视化与 confidence threshold")
    paragraph(doc, "visualize_yolo_predictions.py 在同一 validation sample（sample index=0）上并列展示 RGB、crop+weed r010 ground truth/prediction、weed-only r010 ground truth/prediction。crop+weed r010 的预测框相对更克制，weed-only r010 更容易产生密集候选框或重复框。该观察仅来自单张样本，不能替代整体验证集指标。")
    paragraph(doc, "将 confidence threshold 从 0.25 提高到 0.40 后，低置信度框明显减少，展示更干净；但提高阈值也可能增加漏检，不能把 0.40 直接解释为最佳部署阈值。confidence filtering 与 NMS 均属于推理后处理，不改变 backbone、训练结果或数据集 bbox 过滤规则。")

    section(doc, 17, "YOLO Post-processing Threshold Analysis")
    paragraph(doc, "分析对象为当前最佳 YOLO baseline：YOLOv8n crop+weed r010。使用前 20 张 validation images，比较不同 confidence threshold 和 NMS IoU threshold 对预测框数量、weed 预测框数量及平均置信度的影响。confidence filtering 与 NMS 属于推理后处理，不改变模型训练或数据集 bbox 过滤规则。")
    doc.add_heading("Confidence threshold 是主要影响因素", level=2)
    paragraph(doc, "固定 NMS IoU = 0.50，结果如下。crop 和 weed 框数均为 20 张图像的合计。")
    table(doc, ["Confidence", "Mean boxes/image", "Crop boxes", "Weed boxes", "Mean confidence"], [
        ("0.25", "34.25", "263", "422", "0.4140"),
        ("0.30", "25.15", "179", "324", "0.4644"),
        ("0.40", "14.40", "95", "193", "0.5539"),
        ("0.50", "8.60", "55", "117", "0.6243"),
    ], [2.5, 4.2, 3.0, 3.0, 4.3], 8.2)
    paragraph(doc, "从 conf=0.25 提高到 conf=0.40 后，mean boxes/image 从 34.25 降到 14.40，weed predicted boxes 从 422 降到 193，可视化结果明显更干净。")
    doc.add_heading("NMS IoU threshold 影响较小", level=2)
    paragraph(doc, "固定 conf=0.40，结果如下。")
    table(doc, ["NMS IoU", "Mean boxes/image", "Weed boxes"], [
        ("0.50", "14.40", "193"),
        ("0.60", "14.55", "195"),
        ("0.70", "14.65", "196"),
    ], [4.0, 7.0, 6.0], 8.8)
    paragraph(doc, "在当前前 20 张 validation images 上，NMS IoU threshold 对预测框数量影响较小，而 confidence threshold 更关键。conf=0.40, iou=0.50 可以作为当前 YOLO prediction visualization 的候选设置：它比 conf=0.25 更干净，同时不像 conf=0.50 那样过度减少预测框。")
    paragraph(doc, "限制：前 20 张图分析只统计预测框数量和平均置信度，没有直接计算 TP、FP、FN，因此不能说明 conf=0.40 是最终最优阈值。")
    doc.add_heading("完整验证集 PR F1 sweep", level=2)
    paragraph(doc, "r010 best checkpoint 的平均 F1 在 conf≈0.164 达到最高值 0.544；crop 与 weed 的最佳 F1 阈值分别约为 0.164 和 0.169。conf=0.25/0.30/0.40/0.50 的平均 F1 分别为 0.509、0.466、0.364、0.247。conf=0.40 的平均 precision 约 0.747，但平均 recall 仅约 0.241，因此只适合作为干净展示候选；需要平衡 precision/recall 时应从 conf≈0.16 开始，再结合误喷与漏喷成本校准。当前 CPU 完整验证 inference 约为 10.4 ms/image。")

    section(doc, 18, "YOLO lightweight baseline 与 MobileNetV3 优化")
    doc.add_heading("YOLOv8n lightweight baseline metrics", level=2)
    paragraph(doc, "下表汇总四个 5 epochs runs。检测指标取 results.csv 最后一轮 all-class 值，Params、GFLOPs（imgsz=480）和 model size 取 best.pt；crop+weed 与 weed-only 的类别口径不同，不宜直接用 all-class 指标比较检测效果。")
    table(doc, ["Run", "类别", "P", "R", "mAP50", "mAP50-95", "Params", "GFLOPs", "Size"], [
        ("r020", "crop+weed", "0.49736", "0.59299", "0.52611", "0.29074", "3,011,238", "4.6084", "6.214 MB"),
        ("r015", "crop+weed", "0.50758", "0.58699", "0.53032", "0.29651", "3,011,238", "4.6084", "6.214 MB"),
        ("r010", "crop+weed", "0.50539", "0.58950", "0.53062", "0.29557", "3,011,238", "4.6084", "6.214 MB"),
        ("weed-only r010", "weed-only", "0.43149", "0.48721", "0.42087", "0.18213", "3,011,043", "4.6078", "6.214 MB"),
    ], [2.6, 1.9, 1.4, 1.4, 1.7, 2.0, 2.1, 1.6, 2.0], 6.8)
    paragraph(doc, "当前最佳 YOLO baseline 是 crop+weed r010：mAP50=0.53062、mAP50-95=0.29557、约 3.01M Params、4.61 GFLOPs、6.21 MB。")
    doc.add_heading("MobileNetV3 Small YOLOv8 optimization workflow", level=2)
    paragraph(doc, "已实现 MobileNetV3-Small backbone，输出 P3/8、P4/16、P5/32 特征，后接 YOLOv8-style PAN-FPN 与 anchor-free Detect head。backbone 加载官方 ImageNet 预训练权重，模型入口显式完成 ImageNet mean/std 归一化。训练采用两阶段方案：冻结 backbone 5 epochs 预热新 neck/head，再从最佳 warm-up 权重解冻全模型训练 20 epochs。")
    table(doc, ["模型", "P", "R", "mAP50", "mAP50-95", "Params", "GFLOPs", "Size", "CPU"], [
        ("YOLOv8n r010", "0.5049", "0.5909", "0.5306", "0.2956", "3.01M", "4.61", "6.21 MB", "10.4 ms"),
        ("MobileNetV3 初始 5 ep", "0.2251", "0.2594", "0.1284", "0.0325", "2.39M", "2.55", "5.04 MB", "25–41 ms"),
        ("MobileNetV3 两阶段", "0.4773", "0.5269", "0.4810", "0.2106", "2.39M", "2.55", "5.04 MB", "40.8 ms"),
    ], [4.0, 1.2, 1.2, 1.5, 1.8, 1.8, 1.3, 1.7, 2.1], 6.6)
    paragraph(doc, "修正归一化并完成两阶段训练后，mAP50 从 0.1284 恢复到 0.4810。相对 YOLOv8n r010，参数量、GFLOPs 和模型文件分别减少约 20.6%、44.7%、18.9%，但 mAP50 与 mAP50-95 仍分别低 0.0496 和 0.0850；当前 CPU inference 40.8 ms/image，约为 baseline 的 3.9 倍。轻量化体积和理论计算量目标已达到，但精度与当前 CPU latency 尚未超过 baseline，因此暂不替换 YOLOv8n r010。后续应在目标硬件重新测速，并继续研究 neck/head 与高 IoU 定位精度。")
    paragraph(doc, "检测支线已完成 detection dataset 构建、bbox statistics、r020/r015/r010 过滤对比、weed-only 对照、prediction visualization、后处理阈值分析、lightweight baseline 汇总和 MobileNetV3 两阶段优化。由于 common split 以 dense vegetation 为主，项目主线仍是 MobileNetV2ShallowUNet + boundary weighted CE r5_w4 语义分割；YOLO 作为 sparse 场景候选路线。")

    section(doc, 19, "YOLO / U-Net 路线选择标准")
    paragraph(doc, "无人机多光谱图像 → NDVI / 植物-土壤区分指标 → 计算 VCR → 判断 sparse / transition / dense → 选择定位或区域分割路线。")
    table(doc, ["VCR 场景", "候选路线", "除草用途"], [
        ("sparse", "YOLO", "单株或单簇定位，点状精准除草"),
        ("transition", "同时测试 YOLO 与 U-Net，或人工确认", "根据目标形态与作业要求选择"),
        ("dense", "U-Net / MobileNetV2ShallowUNet", "输出区域 mask，支持区域除草"),
    ], [3, 6.8, 7.2], 8.6)
    paragraph(doc, "common split 中 dense 占约 95.15%，因此当前阶段以 U-Net / MobileNetV2ShallowUNet 语义分割为主路线有数据依据。YOLO 已完成初步检测对照，适合作为低覆盖稀疏场景的候选路线或后续扩展；尚无与语义分割模型在相同条件下的正式效果对比。")

    section(doc, 20, "给老师汇报用总结")
    paragraph(doc, "老师，目前我已经完成了从模拟数据到真实 WeedMap 多光谱无人机数据的完整语义分割流程。前期我先用 NDVI、NDRE 做了作物、杂草和土壤的光谱差异模拟，之后用 synthetic 数据训练 U-Net，发现普通 CE 会忽视 weed，加入 class weights 后 weed IoU 大幅提升。接着我整理了真实 WeedMap Tiles 数据并完成标签映射验证；本地数据中 color 标签的类别语义最明确，因此 Dataset 从 GroundTruth_color 生成 0/1/2 标签，并把 mask=255 作为 ignore 区域。严格公平的 RGB 与 multispectral 对比使用相同的 454 个样本，其中 train 363 个、val 91 个。Multispectral 的 mean IoU 和 weed IoU 均高于 RGB，说明多光谱通道对杂草识别更有帮助。")
    paragraph(doc, "在真实数据实验中，Weighted CE 是最稳定的基础 loss。进一步分析 sample index=0 后发现，5px 边界区域包含 Weighted CE 的 95.18% 错误，因此我加入了 boundary weighted CE。SmallUNet 上 boundary weighted CE r5_w4 的三 seed 平均 Mean IoU 达到 75.17%，Weed IoU 达到 58.29%。之后我实现了 MobileNetV2ShallowUNet 方案 A，它不是论文完整 MobileNetV2-U-Net 复现，而是浅层 MobileNetV2-style encoder 改造，保留当前两级 U-Net decoder。该模型在 weighted CE 下已经优于原 SmallUNet。进一步将 MobileNetV2ShallowUNet 与 boundary weighted CE r5_w4 结合后，三 seed 平均 Mean IoU 达到 76.71% ± 0.55%，Weed IoU 达到 59.82% ± 0.64%，是目前最好的结果。说明浅层 MobileNetV2-style encoder 与 boundary-aware loss 可以叠加提升，尤其对 weed 类识别更有帮助。")
    paragraph(doc, "另外，我新增了一个基于植被覆盖率 VCR 的模型选择指标。通过 NDVI 区分植物与土壤，并计算每张图像的植被覆盖率。统计结果显示 WeedMap common split 中 dense 样本占 95.15%，说明当前数据集以密集植被覆盖场景为主，因此继续优化 U-Net / MobileNetV2ShallowUNet 这类语义分割模型是合理的；YOLO 更适合作为低覆盖稀疏场景下的补充路线，用于单株或单簇定位和点状精准除草。")
    paragraph(doc, "VCR regression attention router 已完成 Tiny CNN、SE-Tiny CNN、CBAM-Tiny CNN 的 30 epochs × 3 seeds 正式实验。CBAM-Tiny CNN 的 MAE 最低，为 0.0923 ± 0.0133，但 sparse recall 仅 0.1111 ± 0.1925。该实验验证了从原始 G/R/RE/NIR 回归连续 VCR 的初步可行性，但受 non-dense 样本极少和空间数据泄漏风险限制，当前不能替代 NDVI-VCR 规则，也不能部署；它用于未来连接 segmentation 与 detection 两条分支。")
    paragraph(doc, "YOLO 检测支线已完成 mask-to-bbox 数据构建、bbox 统计、r020/r015/r010 过滤对比、weed-only 对照、prediction visualization、后处理阈值分析和 lightweight baseline 汇总。当前最佳 YOLOv8n crop+weed r010 的 mAP50 为 0.53062、mAP50-95 为 0.29557，约 3.01M Params、4.61 GFLOPs、6.21 MB。前 20 张图的框数分析显示 confidence threshold 比 NMS IoU 更影响输出数量；完整验证集平均 F1 的最佳起点约为 conf=0.16，conf=0.40 只适合较干净的展示。MobileNetV3-Small-YOLOv8 经归一化修正和 5+20 epochs 两阶段训练后达到 mAP50=0.4810、mAP50-95=0.2106，参数量和 GFLOPs 分别减少约 20.6% 和 44.7%，但精度仍低、当前 CPU inference 反而更慢，因此不替换 r010 baseline。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Updated {OUTPUT}")


if __name__ == "__main__":
    build()
